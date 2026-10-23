"""Throwaway converter: the report Markdown -> styled .docx (scratch, not part of the app)."""
import re
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x6B, 0x72, 0x80)
AMBER = RGBColor(0xB4, 0x6A, 0x00)
NAVY = RGBColor(0x1D, 0x4E, 0xD8)
CODE_BG = "F4F1EA"
TABLE_BG = "FAF7F2"

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def shade(par_el, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    par_el.get_or_add_pPr().append(shd)


def shade_cell(cell, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def add_runs(par, text, base_bold=False, color=None, size=None):
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        bold, code, ital, link = False, False, False, None
        body = chunk
        if chunk.startswith("**") and chunk.endswith("**"):
            body, bold = chunk[2:-2], True
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            body, ital = chunk[1:-1], True
        elif chunk.startswith("`") and chunk.endswith("`"):
            body, code = chunk[1:-1], True
        elif chunk.startswith("["):
            m = LINK.fullmatch(chunk)
            if m:
                body, link = m.group(1), m.group(2)
        run = par.add_run(body)
        run.bold = bold or base_bold
        run.italic = ital
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt((size or 10.5) - 0.5)
            run.font.color.rgb = AMBER
        if link:
            run.font.color.rgb = NAVY
            run.underline = True
            _hyperlink(par, link, run)
        elif color is not None and not code:
            run.font.color.rgb = color
        if size and not code:
            run.font.size = Pt(size)


def _hyperlink(par, url, run):
    """Move an already-built run inside a real w:hyperlink element."""
    r = run._r
    r.getparent().remove(r)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), par.part.relate_to(url, RT.HYPERLINK, is_external=True))
    link.append(r)
    par._p.append(link)


def heading(doc, text, level):
    sizes = {1: 15, 2: 12.5, 3: 11}
    par = doc.add_paragraph(style=f"Heading {level}")
    par.paragraph_format.space_before = Pt(16 if level == 1 else 14)
    par.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    par.paragraph_format.keep_with_next = True
    add_runs(par, text, base_bold=True, color=AMBER if level == 1 else INK, size=sizes[level])
    if level == 1:
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "ECE5DB")
        border.append(bottom)
        par._p.get_or_add_pPr().append(border)
    return par


def code_block(doc, lines):
    for i, line in enumerate(lines):
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(6 if i == 0 else 0)
        par.paragraph_format.space_after = Pt(6 if i == len(lines) - 1 else 0)
        par.paragraph_format.left_indent = Cm(0.3)
        par.paragraph_format.right_indent = Cm(0.3)
        run = par.add_run(line.replace("\t", "    ") or " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = INK
        shade(par._p, CODE_BG)


def table(doc, rows):
    header, *body = rows
    empty_header = not any(c.strip() for c in header)
    if empty_header:
        rows = body
    else:
        rows = [header] + body
    if not rows:
        return
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for r_idx, row in enumerate(rows):
        cells = t.add_row().cells
        for c_idx in range(cols):
            text = row[c_idx] if c_idx < len(row) else ""
            par = cells[c_idx].paragraphs[0]
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.space_after = Pt(2)
            is_head = (not empty_header) and r_idx == 0
            add_runs(par, text, base_bold=is_head, size=9.5)
            if is_head:
                shade_cell(cells[c_idx], TABLE_BG)
                for run in par.runs:
                    run.font.color.rgb = INK
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def bullet_list(doc, items, numbered):
    style = "List Number" if numbered else "List Bullet"
    for text in items:
        par = doc.add_paragraph(style=style)
        par.paragraph_format.space_after = Pt(2)
        add_runs(par, text)


def parse_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(md_path, docx_path):
    lines = open(md_path, encoding="utf-8").read().splitlines()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(1.9)
        section.left_margin = section.right_margin = Cm(2.1)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("CivicRoute AI  ·  page ")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    i, para, bullets = 0, [], None
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if bullets:
                bullet_list(doc, bullets[1], bullets[0])
                bullets = None
            if para:
                p = doc.add_paragraph()
                add_runs(p, " ".join(para))
                para = []
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            code_block(doc, block)
            i += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            if bullets:
                bullet_list(doc, bullets[1], bullets[0])
                bullets = None
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = parse_table_row(lines[i])
                if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                    rows.append(cells)
                i += 1
            table(doc, rows)
            continue

        if stripped == "---" or stripped == "":
            if para:
                p = doc.add_paragraph()
                add_runs(p, " ".join(para))
                para = []
            if bullets:
                bullet_list(doc, bullets[1], bullets[0])
                bullets = None
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            if para:
                p = doc.add_paragraph()
                add_runs(p, " ".join(para))
                para = []
            if bullets:
                bullet_list(doc, bullets[1], bullets[0])
                bullets = None
            level = min(len(m.group(1)), 3)
            if level == 1 and not doc.paragraphs[1:]:
                par = doc.add_paragraph(style="Title")
                par.paragraph_format.space_after = Pt(2)
                add_runs(par, m.group(2), base_bold=True, color=INK, size=26)
            else:
                heading(doc, m.group(2), level)
            i += 1
            continue

        if stripped.startswith("> "):
            if bullets:
                bullet_list(doc, bullets[1], bullets[0])
                bullets = None
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.7)
            par.paragraph_format.space_before = Pt(4)
            par.paragraph_format.space_after = Pt(8)
            add_runs(par, " ".join(quote), color=NAVY)
            border = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), "F59E0B")
            border.append(left)
            par._p.get_or_add_pPr().append(border)
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            if para:
                p = doc.add_paragraph()
                add_runs(p, " ".join(para))
                para = []
            if bullets and not bullets[0]:
                bullet_list(doc, bullets[1], False)
                bullets = None
            bullets = bullets or [True, []]
            bullets[0] = True
            bullets[1].append(m.group(2))
            i += 1
            continue

        if stripped.startswith("- "):
            if para:
                p = doc.add_paragraph()
                add_runs(p, " ".join(para))
                para = []
            if bullets and bullets[0]:
                bullet_list(doc, bullets[1], True)
                bullets = None
            bullets = bullets or [False, []]
            bullets[0] = False
            bullets[1].append(stripped[2:])
            i += 1
            continue

        if bullets:
            bullets[1][-1] += " " + stripped
            i += 1
            continue
        para.append(stripped)
        i += 1

    if bullets:
        bullet_list(doc, bullets[1], bullets[0])
    if para:
        p = doc.add_paragraph()
        add_runs(p, " ".join(para))

    doc.save(docx_path)
    print("wrote", docx_path)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
